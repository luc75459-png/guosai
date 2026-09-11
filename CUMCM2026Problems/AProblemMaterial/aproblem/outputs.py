# NPZ/CSV 预览与正式 Excel 输出
"""拆分仿真结果，写出预览文件，并生成问题 4 的正式结果工作簿。"""

from __future__ import annotations

import copy
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd
from openpyxl.cell.cell import Cell

from .integrator import SimulationResult
from .model import DryingModel


PHYSICAL_DISTANCE_CM = np.arange(0.0, 2.0, 0.1)
SURFACE_COLUMN_NAME = "药材表面"


def split_result(model: DryingModel, result: SimulationResult) -> tuple[np.ndarray, np.ndarray]:
    """按模型节点数将状态矩阵拆分为温度矩阵和含水率矩阵。"""

    n = model.node_count
    return result.state[:, :n], result.state[:, n:]


def write_preview_files(
    output_dir: Path,
    question: int,
    model: DryingModel,
    result: SimulationResult,
) -> list[Path]:
    """保存全精度压缩结果和带中文表头的 CSV 预览文件。

    问题4的 CSV 使用归一化半径；正式 Excel 导出时还需映射到题目要求的
    固定物理距离。NPZ 保留英文键名，便于后续代码稳定读取。
    """

    output_dir.mkdir(parents=True, exist_ok=True)
    temperature, moisture = split_result(model, result)
    initial_radius = model.radius(0.0)
    distance_cm = model.grid.geometry(initial_radius).nodes_m * 100.0
    radial_fraction = np.linspace(0.0, 1.0, model.node_count)
    radius_m = np.asarray([model.radius(time_s) for time_s in result.time_s])

    prefix = output_dir / f"question{question}"
    npz_path = prefix.with_suffix(".npz")
    np.savez_compressed(
        npz_path,
        time_s=result.time_s,
        temperature_c=temperature,
        moisture=moisture,
        initial_distance_cm=distance_cm,
        radial_fraction=radial_fraction,
        radius_m=radius_m,
        cylinder_length_m=model.cylinder_length_m,
        include_end_faces=model.include_end_faces,
        event_time_s=np.nan if result.event_time_s is None else result.event_time_s,
        step_s=np.asarray([]) if result.step_s is None else result.step_s,
    )

    if question == 4:
        columns = [f"归一化半径 ξ={value:.4f}" for value in radial_fraction]
    else:
        columns = [f"距中心 {value:.4f} cm" for value in distance_cm]
    temperature_frame = pd.DataFrame(temperature, columns=columns)
    temperature_frame.insert(0, "时间（s）", result.time_s)
    moisture_frame = pd.DataFrame(moisture, columns=columns)
    moisture_frame.insert(0, "时间（s）", result.time_s)

    temperature_path = output_dir / f"question{question}_temperature_preview.csv"
    moisture_path = output_dir / f"question{question}_moisture_preview.csv"
    temperature_frame.to_csv(temperature_path, index=False, encoding="utf-8-sig")
    moisture_frame.to_csv(moisture_path, index=False, encoding="utf-8-sig")
    return [npz_path, temperature_path, moisture_path]


def map_material_profile_to_physical(
    field: np.ndarray,
    radius_m: float,
    distances_cm: np.ndarray | None = None,
) -> tuple[np.ndarray, float]:
    """把固定材料坐标上的剖面映射到题目要求的物理距离。

    返回 ``(固定物理距离处的值, 药材表面值)``。物理距离超过当前药材半径
    时对应值为 ``nan``，正式写入 Excel 时转换为空白单元格。
    """

    if radius_m <= 0.0:
        raise ValueError("药材半径必须为正数")
    values = np.asarray(field, dtype=float)
    if values.ndim != 1 or values.size < 2:
        raise ValueError("径向剖面必须是一维数组且至少包含两个节点")

    target_cm = PHYSICAL_DISTANCE_CM if distances_cm is None else np.asarray(
        distances_cm,
        dtype=float,
    )
    material_coordinate = np.linspace(0.0, 1.0, values.size)
    target_coordinate = target_cm * 0.01 / radius_m
    mapped = np.full(target_cm.shape, np.nan, dtype=float)
    inside = target_coordinate <= 1.0 + 1.0e-12
    mapped[inside] = np.interp(
        target_coordinate[inside],
        material_coordinate,
        values,
    )
    return mapped, float(values[-1])


def _round_output_value(value: float) -> float | None:
    """按题目要求保留四位小数；超出药材边界的值返回空白。"""

    if not np.isfinite(value):
        return None
    rounded = round(float(value), 4)
    return 0.0 if rounded == 0.0 else rounded


def _copy_cell_style(source: Cell, target: Cell) -> None:
    """复制模板样式，同时保持边框、字体和对齐方式。"""

    target._style = copy.copy(source._style)
    target.number_format = source.number_format


def _result_rows_for_excel(
    result: SimulationResult,
    save_every_s: float = 60.0,
) -> list[tuple[float, int]]:
    """选取正式结果行：所有整分钟输出，并保留最终事件时刻。"""

    rows: list[tuple[float, int]] = []
    tolerance = 1.0e-7
    for index, time_s in enumerate(result.time_s):
        if time_s < save_every_s - tolerance:
            continue
        quotient = time_s / save_every_s
        if np.isclose(quotient, round(quotient), atol=1.0e-8):
            if not rows or not np.isclose(
                rows[-1][0],
                time_s,
                atol=tolerance,
                rtol=0.0,
            ):
                rows.append((float(time_s), index))

    if result.event_time_s is not None:
        event_time = float(result.event_time_s)
        if not rows or not np.isclose(
            rows[-1][0],
            event_time,
            atol=tolerance,
            rtol=0.0,
        ):
            rows.append((event_time, result.time_s.size - 1))
    return rows


def write_result4_workbook(
    template_path: Path,
    output_path: Path,
    model: DryingModel,
    result: SimulationResult,
) -> Path:
    """基于官方模板生成问题 4 的 ``result4.xlsx``。

    列结构为 ``0.0, 0.1, ..., 1.9 cm`` 加单独的药材表面列。超过当前
    药材半径的位置保持空白，不填零；所有水分值保留四位小数。
    """

    if not template_path.exists():
        raise FileNotFoundError(f"未找到问题 4 结果模板：{template_path}")

    workbook = openpyxl.load_workbook(template_path)
    worksheet = workbook.active
    _, moisture = split_result(model, result)

    base_header_style = copy.copy(worksheet.cell(row=1, column=2)._style)

    for row in range(2, worksheet.max_row + 1):
        for column in range(1, max(worksheet.max_column, 22) + 1):
            worksheet.cell(row=row, column=column).value = None

    worksheet.cell(row=1, column=1, value="时间\\到药材中心的距离")
    for column, distance_cm in enumerate(PHYSICAL_DISTANCE_CM, start=2):
        cell = worksheet.cell(
            row=1,
            column=column,
            value=round(float(distance_cm), 1),
        )
        _copy_cell_style(worksheet.cell(row=1, column=2), cell)
        cell.number_format = "0.0"
    surface_header = worksheet.cell(
        row=1,
        column=len(PHYSICAL_DISTANCE_CM) + 2,
        value=SURFACE_COLUMN_NAME,
    )
    _copy_cell_style(worksheet.cell(row=1, column=2), surface_header)
    worksheet.cell(row=1, column=1)._style = copy.copy(base_header_style)

    excel_rows = _result_rows_for_excel(result)
    for row_index, (time_s, state_index) in enumerate(excel_rows, start=2):
        radius_m = model.radius(time_s)
        physical_values, surface_value = map_material_profile_to_physical(
            moisture[state_index],
            radius_m,
        )

        time_cell = worksheet.cell(row=row_index, column=1, value=round(time_s, 6))
        _copy_cell_style(worksheet.cell(row=2, column=1), time_cell)
        time_cell.number_format = "0.0000"

        for column, value in enumerate(physical_values, start=2):
            cell = worksheet.cell(
                row=row_index,
                column=column,
                value=_round_output_value(value),
            )
            _copy_cell_style(worksheet.cell(row=2, column=column), cell)
            cell.number_format = "0.0000"

        surface_cell = worksheet.cell(
            row=row_index,
            column=len(PHYSICAL_DISTANCE_CM) + 2,
            value=_round_output_value(surface_value),
        )
        _copy_cell_style(worksheet.cell(row=2, column=2), surface_cell)
        surface_cell.number_format = "0.0000"

    for column in range(2, len(PHYSICAL_DISTANCE_CM) + 3):
        worksheet.column_dimensions[
            openpyxl.utils.get_column_letter(column)
        ].width = 9.625

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    workbook.close()
    return output_path
