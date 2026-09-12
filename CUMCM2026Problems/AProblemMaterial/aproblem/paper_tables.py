# -*- coding: utf-8 -*-
"""从官方 resultN.xlsx 抽取论文表 1~6，避免手工取点。

用法：在 AProblemMaterial 目录下执行
``python -m aproblem.paper_tables --question 1``（默认自动定位 resultN.xlsx），
或用 ``--path`` 指定具体工作簿、用 ``--out`` 写出 Markdown 文件。

设计要点：

* 数值一律从 resultN.xlsx **回读**（与交付文件同源），不重新计算、不手抄；
* 选行按最近邻时刻、选列按最近邻距离；网格取 20 的倍数时两者都精确命中；
* 问题4 的表6 末列固定取"药材表面"。
"""

from __future__ import annotations

import argparse
from pathlib import Path

import numpy as np
from openpyxl import load_workbook

from .config import ProjectPaths

RADII_CM = (0.0, 0.5, 1.0, 1.5, 2.0)


def read_sheet(path: Path, sheet_name: str):
    """回读一个工作表：返回 (时间数组, 列头列表, 数值矩阵)。"""

    workbook = load_workbook(path, read_only=True, data_only=True)
    if sheet_name not in workbook.sheetnames:
        raise KeyError(
            f"{path.name} 中没有工作表 {sheet_name}（现有：{workbook.sheetnames}）"
        )
    rows = workbook[sheet_name].iter_rows(values_only=True)
    header = next(rows)
    columns = list(header[1:])
    times: list[float] = []
    values: list[list[float]] = []
    for row in rows:
        if row[0] is None:
            continue
        times.append(float(row[0]))
        values.append([np.nan if item is None else float(item) for item in row[1:]])
    workbook.close()
    return np.asarray(times), columns, np.asarray(values)


def _nearest(array: np.ndarray, target: float) -> int:
    return int(np.argmin(np.abs(array - target)))


def _radius_index(columns, target_cm: float) -> int:
    """在列头里找最接近 target_cm 的距离列。"""

    best_index, best_gap = None, None
    for index, column in enumerate(columns):
        try:
            value = float(column)
        except (TypeError, ValueError):
            continue
        gap = abs(value - target_cm)
        if best_gap is None or gap < best_gap:
            best_index, best_gap = index, gap
    if best_index is None:
        raise ValueError(f"列头里找不到距离列：{columns}")
    return best_index


def _surface_index(columns) -> int:
    """找"药材表面"列；找不到就取最后一列。"""

    for index, column in enumerate(columns):
        if isinstance(column, str) and "表面" in column:
            return index
    return len(columns) - 1


def _format(value: float) -> str:
    return "—" if not np.isfinite(value) else f"{value:.4f}"


def _table(title: str, times_used, matrix, column_labels) -> str:
    lines = [f"**{title}**", ""]
    lines.append("| t | " + " | ".join(column_labels) + " |")
    lines.append("|" + "---|" * (len(column_labels) + 1))
    for label, row in zip(times_used, matrix):
        lines.append(f"| {label} | " + " | ".join(_format(v) for v in row) + " |")
    return "\n".join(lines)


def _time_labels(times_s, hours: bool) -> list[str]:
    if hours:
        return [f"{float(t) / 3600:.2f} h" for t in times_s]
    return [f"{float(t):g} s" for t in times_s]


def build_tables(question: int, path: Path) -> str:
    """按题号生成对应的论文表格（Markdown）。"""

    if question == 1:
        want_s, hours = (100, 300, 600, 900, 1200, 1500, 1800), False
        table_temperature, table_moisture = "表1　药材的温度（℃）", "表2　药材的含水率（kg/kg）"
    elif question == 2:
        want_s, hours = (1800, 3600, 5400, 7200, 9000, 10800), True
        table_temperature, table_moisture = "表3　药材的温度（℃）", "表4　药材的含水率（kg/kg）"
    elif question == 3:
        want_s, hours = None, True
        table_temperature, table_moisture = "", "表5　药材烘干过程的含水率（kg/kg）"
    elif question == 4:
        want_s, hours = None, True
        table_temperature, table_moisture = "", "表6　药材烘干过程的含水率（kg/kg）"
    else:
        raise ValueError("问题编号必须是 1、2、3、4 之一")

    blocks: list[str] = []

    if table_temperature:
        times, columns, values = read_sheet(path, "温度")
        index = [_radius_index(columns, r) for r in RADII_CM]
        chosen = [_nearest(times, t) for t in want_s]
        blocks.append(_table(
            table_temperature,
            _time_labels(times[chosen], hours),
            values[chosen][:, index],
            [f"{r:g} cm" for r in RADII_CM],
        ))

    times, columns, values = read_sheet(path, "水分浓度" if question in (1, 2) else "Sheet1")
    if question in (1, 2):
        index = [_radius_index(columns, r) for r in RADII_CM]
        labels = [f"{r:g} cm" for r in RADII_CM]
        targets = list(want_s)
    else:
        index = [_radius_index(columns, r) for r in RADII_CM[:-1]] + [_surface_index(columns)]
        labels = [f"{r:g} cm" for r in RADII_CM[:-1]] + ["药材表面"]
        step = 6 * 3600.0
        targets = list(np.arange(step, times[-1], step)) + [times[-1]]
    chosen = [_nearest(times, t) for t in targets]
    blocks.append(_table(
        table_moisture,
        _time_labels(times[chosen], hours),
        values[chosen][:, index],
        labels,
    ))

    header = (
        f"# 论文表格（问题{question}）\n\n"
        f"> 源文件：`{path}`　｜　由 `python -m aproblem.paper_tables` 自动抽取，未手工修改\n"
    )
    return header + "\n" + "\n\n".join(blocks) + "\n"


def default_result_path(question: int) -> Path:
    """优先 deliverables，其次 outputs 下最新的一个 resultN.xlsx。"""

    paths = ProjectPaths.discover()
    delivered = paths.project_root.parent / "deliverables" / f"result{question}.xlsx"
    if delivered.exists():
        return delivered
    candidates = sorted(
        (paths.project_root / "outputs").glob(f"问题{question}/**/result{question}.xlsx"),
        key=lambda item: item.stat().st_mtime,
        reverse=True,
    )
    if not candidates:
        raise FileNotFoundError(f"找不到 result{question}.xlsx，请用 --path 指定")
    return candidates[0]


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="从 resultN.xlsx 抽取论文表 1~6")
    parser.add_argument("--question", type=int, choices=(1, 2, 3, 4), required=True)
    parser.add_argument("--path", type=Path, default=None, help="resultN.xlsx 的路径")
    parser.add_argument("--out", type=Path, default=None, help="输出 Markdown 文件（默认打印到屏幕）")
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    path = args.path or default_result_path(args.question)
    text = build_tables(args.question, path)
    if args.out is None:
        print(text)
        return
    args.out.parent.mkdir(parents=True, exist_ok=True)
    args.out.write_text(text, encoding="utf-8")
    print(f"已写出 {args.out}")


if __name__ == "__main__":
    main()
