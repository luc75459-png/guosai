#基础范围检查
"""对仿真结果执行有限性、非负性和物理范围等基础检查。"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

import numpy as np
import openpyxl

from .integrator import SimulationResult
from .model import DryingModel
from .outputs import split_result


@dataclass(frozen=True)
class ValidationReport:
    """汇总一次仿真的基础数值与物理检查结果。"""

    all_finite: bool
    nonnegative_moisture: bool
    material_mesh_consistent: bool
    maximum_radius_error_m: float
    minimum_node_spacing_m: float
    minimum_step_s: float
    maximum_step_s: float
    temperature_bounds_c: tuple[float, float]
    moisture_bounds: tuple[float, float]
    final_max_moisture: float

    def __str__(self) -> str:
        """以中文输出验证摘要，便于终端阅读和记录。"""

        finite_text = "是" if self.all_finite else "否"
        nonnegative_text = "是" if self.nonnegative_moisture else "否"
        mesh_text = "是" if self.material_mesh_consistent else "否"
        return (
            f"数值全部有限：{finite_text}；含水率均非负：{nonnegative_text}；"
            f"材料网格一致：{mesh_text}；"
            f"最大外半径误差：{self.maximum_radius_error_m:.3e} m；"
            f"最小节点间距：{self.minimum_node_spacing_m:.6g} m；"
            f"实际步长范围：{self.minimum_step_s:.6g}～"
            f"{self.maximum_step_s:.6g} s；"
            f"温度范围：{self.temperature_bounds_c[0]:.4f}～"
            f"{self.temperature_bounds_c[1]:.4f} ℃；"
            f"含水率范围：{self.moisture_bounds[0]:.6f}～"
            f"{self.moisture_bounds[1]:.6f} kg/kg；"
            f"最终全场最大含水率：{self.final_max_moisture:.6f} kg/kg"
        )


def validate_result(model: DryingModel, result: SimulationResult) -> ValidationReport:
    """根据温度和含水率矩阵生成基础验证报告。"""

    temperature, moisture = split_result(model, result)
    maximum_radius_error = 0.0
    minimum_spacing = float("inf")
    material_mesh_consistent = True
    expected_fraction = np.linspace(0.0, 1.0, model.node_count)
    for time_s in result.time_s:
        radius_m = model.radius(float(time_s))
        geometry = model.grid.geometry(radius_m)
        maximum_radius_error = max(
            maximum_radius_error,
            abs(float(geometry.nodes_m[-1] - radius_m)),
        )
        spacing = np.diff(geometry.nodes_m)
        minimum_spacing = min(minimum_spacing, float(np.min(spacing)))
        if (
            not np.all(np.isfinite(geometry.nodes_m))
            or np.any(spacing <= 0.0)
            or not np.allclose(
                geometry.nodes_m / radius_m,
                expected_fraction,
                rtol=0.0,
                atol=1.0e-12,
            )
        ):
            material_mesh_consistent = False

    if result.step_s is None or result.step_s.size == 0:
        minimum_step = 0.0
        maximum_step = 0.0
    else:
        minimum_step = float(np.min(result.step_s))
        maximum_step = float(np.max(result.step_s))

    return ValidationReport(
        all_finite=bool(np.all(np.isfinite(result.state))),
        nonnegative_moisture=bool(np.all(moisture >= 0.0)),
        material_mesh_consistent=material_mesh_consistent,
        maximum_radius_error_m=maximum_radius_error,
        minimum_node_spacing_m=minimum_spacing,
        minimum_step_s=minimum_step,
        maximum_step_s=maximum_step,
        temperature_bounds_c=(float(np.min(temperature)), float(np.max(temperature))),
        moisture_bounds=(float(np.min(moisture)), float(np.max(moisture))),
        final_max_moisture=float(np.max(moisture[-1])),
    )


# 问题4结果文件的固定列：0.0～1.9 cm 加一列"药材表面"。
RESULT4_DISTANCE_CM = tuple(round(0.1 * index, 1) for index in range(20))
SURFACE_COLUMN_TITLE = "药材表面"


@dataclass(frozen=True)
class Result4Check:
    """问题4结果工作簿的自动回读校验结果。"""

    sheet_name: str
    data_rows: int
    issues: tuple[str, ...]

    @property
    def passed(self) -> bool:
        """所有检查项是否全部通过。"""

        return not self.issues

    def __str__(self) -> str:
        """输出中文校验摘要。"""

        if self.passed:
            return (
                f"result4 回读校验通过：工作表 {self.sheet_name}，"
                f"数据 {self.data_rows} 行"
            )
        return "result4 回读校验未通过：" + "；".join(self.issues)


def verify_result4_workbook(path: Path) -> Result4Check:
    """回读问题4结果工作簿，逐项检查题目要求的格式。

    检查内容：工作表名、表头（20 个固定距离加"药材表面"）、时间列严格
    递增且为 60 s 的整数倍（末行允许是连续事件时刻）、超出当前半径的
    固定距离单元为空、表面列非空、数值单元格为四位小数格式。
    """

    if not path.exists():
        return Result4Check(path.stem, 0, (f"未找到文件：{path}",))

    workbook = openpyxl.load_workbook(path, data_only=True)
    worksheet = workbook.active
    issues: list[str] = []

    expected_header: list[object] = list(RESULT4_DISTANCE_CM) + [SURFACE_COLUMN_TITLE]
    header = [worksheet.cell(row=1, column=column).value for column in range(2, 23)]
    for index, (actual, expected) in enumerate(zip(header, expected_header), start=2):
        if isinstance(expected, str):
            if actual != expected:
                issues.append(f"第{index}列表头应为{expected}，实际为{actual!r}")
        elif actual is None or not np.isclose(float(actual), expected, atol=1.0e-9):
            issues.append(f"第{index}列表头应为{expected}，实际为{actual!r}")

    if worksheet.max_column > 22:
        # 模板里的省略号占位列必须被清空。
        for column in range(23, worksheet.max_column + 1):
            if worksheet.cell(row=1, column=column).value is not None:
                issues.append(f"第{column}列存在多余表头")

    times: list[float] = []
    for row in range(2, worksheet.max_row + 1):
        value = worksheet.cell(row=row, column=1).value
        if value is None:
            continue
        times.append(float(value))

    if not times:
        issues.append("时间列没有任何数据")
    else:
        if any(later <= earlier for earlier, later in zip(times, times[1:])):
            issues.append("时间列不是严格递增")
        for time_s in times[:-1]:
            if not np.isclose(time_s % 60.0, 0.0, atol=1.0e-6):
                issues.append(f"时间 {time_s} s 不是 60 s 的整数倍")
                break

    for row, time_s in enumerate(times, start=2):
        blanks: list[int] = []
        for column in range(2, 22):
            if worksheet.cell(row=row, column=column).value is None:
                blanks.append(column)
        if blanks and blanks != list(range(blanks[0], 22)):
            issues.append(f"第{row}行的空白单元不连续，超出半径的位置应全部为空")
            break
        if worksheet.cell(row=row, column=22).value is None:
            issues.append(f"第{row}行的药材表面列不得为空")
            break
        for column in range(2, 23):
            cell = worksheet.cell(row=row, column=column)
            if cell.value is None:
                continue
            if cell.number_format != "0.0000":
                issues.append(f"第{row}行第{column}列不是四位小数格式")
                break
            if not np.isfinite(float(cell.value)):
                issues.append(f"第{row}行第{column}列不是有限数")
                break

    sheet_name = worksheet.title
    workbook.close()
    return Result4Check(
        sheet_name=sheet_name,
        data_rows=len(times),
        issues=tuple(issues),
    )


# 问题1~3 官方结果文件的固定列：0.0～2.0 cm，步长 0.1 cm，共 21 列。
RESULT_DISTANCE_CM = tuple(round(0.1 * index, 1) for index in range(21))

# 各题应有的工作表名称；问题3 只有一张表，模板名为 Sheet1。
RESULT_SHEET_TITLES: dict[int, tuple[str, ...] | None] = {
    1: ("温度", "水分浓度"),
    2: ("温度", "水分浓度"),
    3: None,
}

# 各题的首个数据时刻与输出间隔（秒）。
RESULT_FIRST_TIME_S: dict[int, float] = {1: 1.0, 2: 1.0, 3: 60.0}
RESULT_INTERVAL_S: dict[int, float] = {1: 1.0, 2: 1.0, 3: 60.0}


@dataclass(frozen=True)
class ResultWorkbookCheck:
    """问题1~3结果工作簿的回读校验结果。"""

    question: int
    sheets: tuple[str, ...]
    data_rows: int
    issues: tuple[str, ...]

    @property
    def passed(self) -> bool:
        """所有检查项是否全部通过。"""

        return not self.issues

    def __str__(self) -> str:
        """输出中文校验摘要。"""

        if self.passed:
            return (
                f"result{self.question} 回读校验通过：工作表 "
                f"{'、'.join(self.sheets)}，数据 {self.data_rows} 行"
            )
        return (
            f"result{self.question} 回读校验未通过："
            + "；".join(self.issues)
        )


def verify_result_workbook(path: Path, question: int) -> ResultWorkbookCheck:
    """回读问题1~3的结果工作簿，逐项检查题目要求的格式。

    检查内容：工作表名称、表头（0.0～2.0 cm 共 21 列）、时间列严格递增且
    落在规定间隔上（末行允许是连续事件时刻）、全部数据格非空且有限、
    数值单元格为四位小数格式。
    """

    if question not in RESULT_SHEET_TITLES:
        raise ValueError("问题编号必须为 1、2、3 之一")
    if not path.exists():
        return ResultWorkbookCheck(question, (), 0, (f"未找到文件：{path}",))

    workbook = openpyxl.load_workbook(path, data_only=True)
    issues: list[str] = []
    expected_titles = RESULT_SHEET_TITLES[question]
    if expected_titles is not None and tuple(workbook.sheetnames) != expected_titles:
        issues.append(
            f"工作表应为 {expected_titles}，实际为 {tuple(workbook.sheetnames)}"
        )

    rows_seen = 0
    for worksheet in workbook.worksheets:
        header = [
            worksheet.cell(row=1, column=column).value
            for column in range(2, 23)
        ]
        for index, (actual, expected) in enumerate(
            zip(header, RESULT_DISTANCE_CM),
            start=2,
        ):
            if actual is None or not np.isclose(
                float(actual),
                expected,
                atol=1.0e-9,
            ):
                issues.append(
                    f"[{worksheet.title}] 第{index}列表头应为 {expected}，"
                    f"实际为 {actual!r}"
                )

        times: list[float] = []
        for row in range(2, worksheet.max_row + 1):
            value = worksheet.cell(row=row, column=1).value
            if value is None:
                continue
            times.append(float(value))

        if not times:
            issues.append(f"[{worksheet.title}] 时间列没有任何数据")
            continue
        rows_seen = max(rows_seen, len(times))

        if any(later <= earlier for earlier, later in zip(times, times[1:])):
            issues.append(f"[{worksheet.title}] 时间列不是严格递增")
        if not np.isclose(
            times[0],
            RESULT_FIRST_TIME_S[question],
            atol=1.0e-6,
        ):
            issues.append(
                f"[{worksheet.title}] 首个时刻应为 "
                f"{RESULT_FIRST_TIME_S[question]} s，实际为 {times[0]} s"
            )
        interval = RESULT_INTERVAL_S[question]
        for time_s in times[:-1]:
            if not np.isclose(time_s % interval, 0.0, atol=1.0e-6):
                issues.append(
                    f"[{worksheet.title}] 时间 {time_s} s 不是 {interval:g} s "
                    "的整数倍"
                )
                break

        for row in range(2, 2 + len(times)):
            for column in range(2, 23):
                cell = worksheet.cell(row=row, column=column)
                if cell.value is None:
                    issues.append(f"[{worksheet.title}] 第{row}行第{column}列为空")
                    break
                if cell.number_format != "0.0000":
                    issues.append(
                        f"[{worksheet.title}] 第{row}行第{column}列不是四位小数格式"
                    )
                    break
                if not np.isfinite(float(cell.value)):
                    issues.append(
                        f"[{worksheet.title}] 第{row}行第{column}列不是有限数"
                    )
                    break
            else:
                continue
            break

    sheet_names = tuple(workbook.sheetnames)
    workbook.close()
    return ResultWorkbookCheck(
        question=question,
        sheets=sheet_names,
        data_rows=rows_seen,
        issues=tuple(issues),
    )
